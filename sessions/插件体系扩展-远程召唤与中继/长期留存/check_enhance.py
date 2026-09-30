# -*- coding: utf-8 -*-
import openpyxl
SRC = r'D:\TraeWorkSpace\Projects\mecha-sequence-combat\src\plugin-tables\enhance-plugin-tables-v0.9.xlsx'
LO  = r'D:\TraeWorkSpace\Projects\mecha-sequence-combat\tmp\lo-check\enhance-plugin-tables-v0.9.xlsx'
ERRS = ('#REF!', '#NAME?', '#VALUE!', '#DIV/0!', '#N/A', '#NULL!', '#NUM!')

def scan(p, tag):
    wb = openpyxl.load_workbook(p, data_only=True)
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and any(e in c.value for e in ERRS):
                    bad.append((ws.title, c.coordinate, c.value))
    print(f'[{tag}] sheets={wb.sheetnames}')
    print(f'[{tag}] formula-errors={len(bad)} {bad}')

scan(SRC, 'SRC')
scan(LO, 'LO ')

# 主表条数与序号连续性
wb = openpyxl.load_workbook(SRC)
ws = wb['A-赋能主表']
ids, names, tags = [], [], []
for r in range(4, ws.max_row + 1):
    if ws.cell(r, 1).value is None:
        break
    ids.append(ws.cell(r, 1).value)
    names.append(ws.cell(r, 2).value)
    tags.append(ws.cell(r, 4).value)
print('主表条数 =', len(ids), '| 序号连续 =', ids == list(range(1, len(ids) + 1)))
print('通用名 =', ' / '.join(names))
print('Tag 数 =', len(tags), '| 唯一 =', len(set(tags)))

# Tag 体系 / 受附要求体系 的 Tag 覆盖校验
alltags = set()
ws2 = wb['Tag体系']
for r in range(4, ws2.max_row + 1):
    v = ws2.cell(r, 2).value
    if v:
        for t in v.split('/'):
            alltags.add(t.strip())
ws3 = wb['受附要求体系']
reftags = set()
for r in range(4, ws3.max_row + 1):
    v = ws3.cell(r, 1).value
    if v:
        reftags.add(v.strip())
print('Tag体系 覆盖数 =', len(alltags))
print('受附要求体系 覆盖数 =', len(reftags))
print('主表Tag ⊆ Tag体系 :', set(tags) <= alltags)
print('主表Tag ⊆ 受附要求体系 :', set(tags) <= reftags)
print('Tag体系 == 受附要求体系 :', alltags == reftags)
print('缺失(受附要求无) :', set(tags) - reftags)